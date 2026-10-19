import csv
import io
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from app.models.entities import Recipient
from app.schemas.api import RecipientInput

def parse_import(content: bytes, filename: str):
    if filename.lower().endswith('.csv'):
        text = content.decode('utf-8-sig')
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.DictReader(io.StringIO(text), dialect=dialect))
    elif filename.lower().endswith('.xlsx'):
        import zipfile
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(f.file_size for f in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError('Le classeur décompressé dépasse 40 Mo.')
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            iterator = workbook.active.iter_rows(values_only=True)
            headers = [str(x or '').strip().lower() for x in next(iterator)]
            rows = []
            for row in iterator:
                if len(rows) >= 10000:
                    raise ValueError('Maximum 10 000 lignes par import.')
                rows.append(dict(zip(headers, row)))
        finally:
            workbook.close()
    else:
        raise ValueError('Utilisez un fichier CSV UTF-8 ou XLSX.')
    if len(rows) > 10000:
        raise ValueError('Maximum 10 000 lignes par import.')
    if not rows or 'email' not in {str(k).strip().lower() for k in rows[0]}:
        raise ValueError('La colonne email est obligatoire.')
    return rows

def import_recipients(db, content, filename):
    rows = parse_import(content, filename)
    existing = set(db.scalars(select(Recipient.email)))
    report = {'imported': 0, 'duplicates': 0, 'errors': []}
    for line, row in enumerate(rows, 2):
        row = {str(k).strip().lower(): str(v or '').strip() for k, v in row.items() if k}
        try:
            item = RecipientInput(**{k: row.get(k, '') for k in ['email', 'entreprise', 'nom_recruteur', 'poste']})
            if item.email in existing:
                report['duplicates'] += 1
                continue
            db.add(Recipient(**item.model_dump()))
            existing.add(item.email)
            report['imported'] += 1
        except (ValueError, ValidationError):
            report['errors'].append({'line': line, 'message': 'Adresse ou informations invalides.'})
    db.commit()
    return report
